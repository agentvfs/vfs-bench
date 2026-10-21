"""An isolated Mirage Sheets fixture beside the original commerce service."""
import gzip
import hashlib
import io
import json
import os
import shlex
import sqlite3
import subprocess
import tarfile
import tempfile
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils.datetime import to_excel

from server.clients import gws
from server.fakes.sheet_formats import ACCOUNTING_FORMAT, COMPACT_FORMAT
from server.seed.kubernetes import command

REPO = Path(__file__).resolve().parents[2]
REMOTE = '/opt/fixture-api'
INTEG = REMOTE + '/submodule/mirage/integ'
REDIRECT = ('env PYTHONPATH=/opt/fixture-redirect '
            'GWS_MOCK_URL=http://127.0.0.1:19999 ')


def source_folder(task):
    return REPO / 'task/toolathlon' / task / 'shared/fixtures/sheets'


def start(box, task, stack):
    source = json.loads((source_folder(task) / 'source.json').read_text())
    path = source_folder(task) / 'source.xlsx'
    if hashlib.sha256(path.read_bytes()).hexdigest() != source['sha256']:
        raise ValueError('Original workbook checksum mismatch')
    runtime = start_runtime(box, stack)
    seed(task)
    return {'source': source, **runtime}


def mirage_source(*paths):
    """Mirage's own files as the folder has them, not as HEAD committed them.

    Mirage is this project's code and not an upstream this repo pins, so a
    fix made in `submodule/mirage` is meant to be the code that runs.
    Reading the committed blob instead would ship the bytes from before the
    fix and report every hash as matching. `git ls-files` does the
    enumeration so an ignored tree like `node_modules` stays out and a file
    the fix adds comes in.
    """
    mirage = REPO / 'submodule/mirage'
    names = subprocess.check_output([
        'git', '-C',
        str(mirage), 'ls-files', '--cached', '--others', '--exclude-standard',
        '--', *paths
    ],
                                    text=True).splitlines()
    return {
        name: (mirage / name).read_bytes()
        for name in names if (mirage / name).is_file()
    }


def start_runtime(box,
                  stack,
                  wrapper='server/fakes/gws_api.mjs',
                  redirect_path='redirect/eval_shim/gws_redirect.py'):
    """Start the shared Google API transport beside a service fixture."""
    from server.seed.business import expose
    files = {
        'submodule/mirage/' + name: data
        for name, data in mirage_source('integ/server/kit', 'integ/server/gws',
                                        'integ/prisma/gws.prisma',
                                        'integ/fixtures/gws/v1.json').items()
    }
    for name in ('package.json', 'package-lock.json'):
        files['submodule/mirage/integ/' + name] = (REPO / 'docker' /
                                                   name).read_bytes()
    # The proxy borrows its server loop from execution.py, which imports
    # nothing outside the standard library.
    for name in ('server/fakes/sheet_formats.py', 'server/execution.py'):
        files[name] = (REPO / name).read_bytes()
    files[wrapper] = (REPO / wrapper).read_bytes()
    bundle = io.BytesIO()
    with tarfile.open(fileobj=bundle, mode='w:gz') as archive:
        for name, data in files.items():
            entry = tarfile.TarInfo(name)
            entry.size = len(data)
            archive.addfile(entry, io.BytesIO(data))
    box.files.write('/tmp/gws-source.tar.gz', bundle.getvalue(), user='root')
    command(
        box, 'sudo mkdir -p ' + REMOTE + ' && sudo tar -xzf '
        '/tmp/gws-source.tar.gz -C ' + REMOTE)
    command(box,
            'cd ' + INTEG + ' && sudo npm ci --no-audit --no-fund && '
            'sudo env INTEG_DB_URL=file:/tmp/generate.db '
            './node_modules/.bin/prisma generate --schema prisma/gws.prisma',
            timeout=240)
    box.commands.run('cd ' + INTEG + ' && sudo ./node_modules/.bin/tsx ' +
                     REMOTE + '/' + wrapper + ' --port 19998 '
                     '> /tmp/toolathlon-gws.log 2>&1',
                     background=True,
                     timeout=0)
    command(
        box, 'for i in $(seq 1 60); do test -f /tmp/toolathlon-gws-database '
        '&& curl -fsS http://127.0.0.1:19998/_kit/health '
        '>/dev/null && exit 0; '
        'sleep 1; done; cat /tmp/toolathlon-gws.log; exit 1')
    program = (
        'import threading; from server.fakes.sheet_formats import serve; '
        'ctx=serve("http://127.0.0.1:19998", port=19999, '
        'state_path="/tmp/toolathlon-sheet-formats.json"); '
        'ctx.__enter__(); threading.Event().wait()')
    box.commands.run('cd ' + REMOTE + ' && sudo python3 -c ' +
                     shlex.quote(program) +
                     ' > /tmp/toolathlon-formats.log 2>&1',
                     background=True,
                     timeout=0)
    command(
        box, 'for i in $(seq 1 30); do curl -fsS '
        'http://127.0.0.1:19999/_kit/health >/dev/null && exit 0; '
        'sleep 1; done; cat /tmp/toolathlon-formats.log; exit 1')
    port = expose(box, 19999, 0, 8088, stack)
    os.environ['GWS_URL'] = f'http://127.0.0.1:{port}'
    redirect = (REPO / redirect_path).read_bytes()
    box.files.write('/opt/fixture-redirect/sitecustomize.py',
                    redirect,
                    user='root')
    credentials = {
        'token':
        gws.TOKEN,
        'refresh_token':
        gws.TOKEN,
        'client_id':
        'fixture',
        'client_secret':
        'fixture',
        'token_uri':
        'http://127.0.0.1:19999/token',
        'scopes': [
            'https://www.googleapis.com/auth/drive',
            'https://www.googleapis.com/auth/spreadsheets'
        ],
    }
    box.files.write('/opt/toolathlon/configs/google_credentials.json',
                    json.dumps(credentials),
                    user='root')
    return {
        'inputs_sha256': {
            name: hashlib.sha256(data).hexdigest()
            for name, data in files.items()
        },
        'redirect_sha256': hashlib.sha256(redirect).hexdigest()
    }


def seed(task):
    """Import only the workbook that the original preprocessor copies."""
    source = json.loads((source_folder(task) / 'source.json').read_text())
    workbook = openpyxl.load_workbook(source_folder(task) / 'source.xlsx')
    try:
        sheet_id = source['sheet_id']
        gws.call('POST', '/drive/v3/files', {
            'id': sheet_id,
            'name': source['title'],
            'mimeType': gws.SHEET_MIME
        })
        tabs = gws.sheets(sheet_id)
        gws.batch_update(sheet_id, [{
            'updateSheetProperties': {
                'properties': {
                    'sheetId': tabs[0]['properties']['sheetId'],
                    'title': workbook.sheetnames[0]
                },
                'fields': 'title'
            }
        }] + [{
            'addSheet': {
                'properties': {
                    'title': name
                }
            }
        } for name in workbook.sheetnames[1:]])
        for tab in workbook:
            rows = []
            for row in tab.values:
                rows.append([
                    to_excel(value, workbook.epoch) if isinstance(
                        value, (datetime,
                                date)) else value if value is not None else ''
                    for value in row
                ])
            tabs = gws.sheets(sheet_id)
            tab_id = next(t['properties']['sheetId'] for t in tabs
                          if t['properties']['title'] == tab.title)
            grid = next(t['properties']['gridProperties'] for t in tabs
                        if t['properties']['sheetId'] == tab_id)
            resize = []
            for dimension, key, count in [('ROWS', 'rowCount', tab.max_row),
                                          ('COLUMNS', 'columnCount',
                                           tab.max_column)]:
                if count > grid[key]:
                    resize.append({
                        'appendDimension': {
                            'sheetId': tab_id,
                            'dimension': dimension,
                            'length': count - grid[key]
                        }
                    })
                elif count < grid[key]:
                    resize.append({
                        'deleteDimension': {
                            'range': {
                                'sheetId': tab_id,
                                'dimension': dimension,
                                'startIndex': count,
                                'endIndex': grid[key]
                            }
                        }
                    })
            if resize:
                gws.batch_update(sheet_id, resize)
            for row in rows:
                while row and row[-1] == '':
                    row.pop()
            while rows and not rows[-1]:
                rows.pop()
            for offset in range(0, len(rows), 100):
                gws.write_values(sheet_id, f"'{tab.title}'!A{offset + 1}",
                                 rows[offset:offset + 100])
            for offset in range(0, tab.max_row, 100):
                gws.batch_update(sheet_id, [{
                    'updateCells': {
                        'start': {
                            'sheetId': tab_id,
                            'rowIndex': offset,
                            'columnIndex': 0
                        },
                        'rows': [{
                            'values': [{
                                'userEnteredFormat': cell_format(cell)
                            } for cell in row]
                        } for row in tab.iter_rows(
                            min_row=offset + 1,
                            max_row=min(offset + 100, tab.max_row))],
                        'fields':
                        'userEnteredFormat'
                    }
                }])

    finally:
        workbook.close()


def database(box, operation, data=None):
    """Use SQLite's backup API while no task requests are in flight."""
    target = '/tmp/toolathlon-gws-snapshot.db'
    if operation == 'restore':
        box.files.write(target,
                        current_snapshot(gzip.decompress(data)),
                        user='root')
    elif operation != 'capture':
        raise ValueError('Unknown Sheets database operation')
    script = (
        'import sqlite3,pathlib\n'
        "live=pathlib.Path('/tmp/toolathlon-gws-database').read_text()\n"
        f"snapshot={target!r}\n"
        'source,destination=' +
        ('(snapshot,live)' if operation == 'restore' else '(live,snapshot)') +
        '\n'
        'with sqlite3.connect(source) as src, '
        'sqlite3.connect(destination) as dst:\n'
        ' src.backup(dst)\n')
    command(box, 'sudo python3 -c ' + shlex.quote(script))
    if operation == 'capture':
        return gzip.compress(box.files.read(target, format='bytes'), mtime=0)


def current_snapshot(data):
    """Adapt captured Sheets rows to Mirage's current additive schema.

    The committed SQLite capture stays byte-for-byte upstream-derived. Both
    existing commerce captures have empty row/column pixel metadata, so the
    current metadata fields can be initialized without changing their cells.
    """
    with tempfile.NamedTemporaryFile(suffix='.db') as snapshot:
        snapshot.write(data)
        snapshot.flush()
        with sqlite3.connect(snapshot.name) as db:
            tab_columns = {
                row[1]
                for row in db.execute('PRAGMA table_info(SheetTab)')
            }
            cell_columns = {
                row[1]
                for row in db.execute('PRAGMA table_info(SheetCell)')
            }
            if 'rowMeta' not in tab_columns:
                if not {'rowPixels', 'columnPixels'} <= tab_columns:
                    raise ValueError('Unknown captured SheetTab schema')
                if db.execute("SELECT 1 FROM SheetTab WHERE rowPixels != '{}' "
                              "OR columnPixels != '{}' LIMIT 1").fetchone():
                    raise ValueError('Captured Sheets pixel metadata is not '
                                     'empty; cannot migrate it as defaults')
                for column, default in (('rowMeta',
                                         '{}'), ('columnMeta',
                                                 '{}'), ('bandedRanges', '[]'),
                                        ('conditionalFormats', '[]')):
                    db.execute(f'ALTER TABLE SheetTab ADD COLUMN {column} '
                               f'TEXT NOT NULL DEFAULT {default!r}')
                db.execute('ALTER TABLE SheetTab ADD COLUMN basicFilter TEXT')
            if 'props' not in cell_columns:
                db.execute("ALTER TABLE SheetCell ADD COLUMN props TEXT "
                           "NOT NULL DEFAULT '{}'")
            # Drive's change feed is a table of its own, in the DDL Prisma
            # pushes for it. A capture from before it recorded no changes,
            # so it starts empty, as a fresh fake's does.
            if not db.execute(
                    "SELECT 1 FROM sqlite_master WHERE "
                    "type = 'table' AND name = 'DriveChange'").fetchone():
                db.execute('CREATE TABLE "DriveChange" ("tenant" TEXT NOT '
                           'NULL, "seq" INTEGER NOT NULL, "fileId" TEXT NOT '
                           'NULL, PRIMARY KEY ("tenant", "seq"))')
        return Path(snapshot.name).read_bytes()


def state():
    files = []
    cursor = None
    while True:
        path = '/drive/v3/files?pageSize=1000'
        if cursor:
            path += '&pageToken=' + gws.quote(cursor)
        response = gws.call('GET', path)
        files.extend(response['files'])
        cursor = response.get('nextPageToken')
        if not cursor:
            break
    return {
        'files': files,
        'sheets': {
            row['id']: gws.grid(row['id'])
            for row in files if row['mimeType'] == gws.SHEET_MIME
        }
    }


def matches_capture(actual, expected):
    """Compare captured content across Mirage's new presentation defaults."""
    for sheet_id, sheet in actual['sheets'].items():
        captured = expected['sheets'].get(sheet_id, {})
        properties = sheet['properties']
        for field in ('defaultFormat', 'spreadsheetTheme'):
            if field not in captured.get('properties', {}):
                properties.pop(field, None)
    return actual == expected


def set_folder(task):
    folders = gws.find(task, gws.FOLDER_MIME)
    if len(folders) != 1:
        raise ValueError(
            'Original task did not create exactly one Sheets folder')
    os.environ['DRIVE_FOLDER_ID'] = folders[0]['id']


def cell_format(cell):
    """Translate the formats present in the captured original workbooks."""
    if cell.value is None and not cell.has_style and cell.hyperlink is None:
        return {}
    result = {
        'textFormat': {
            'fontFamily': cell.font.name,
            'bold': bool(cell.font.b),
            'italic': bool(cell.font.i)
        }
    }
    if cell.hyperlink and cell.hyperlink.target:
        result['textFormat']['link'] = {'uri': cell.hyperlink.target}
    if cell.font.sz:
        result['textFormat']['fontSize'] = int(cell.font.sz)
    for name, value in [('horizontalAlignment', cell.alignment.horizontal),
                        ('verticalAlignment', cell.alignment.vertical)]:
        if value:
            result[name] = {
                'center':
                'CENTER' if name.startswith('horizontal') else 'MIDDLE'
            }.get(value, value.upper())
    borders = {}
    for edge in ('top', 'bottom', 'left', 'right'):
        border = getattr(cell.border, edge)
        if border is None or border.style is None:
            continue
        if (border.style != 'thin' or border.color.type != 'rgb'
                or border.color.rgb != 'FF000000'):
            raise ValueError('Unsupported source workbook border')
        borders[edge] = {
            'style': 'SOLID',
            'color': {
                'red': 0,
                'green': 0,
                'blue': 0
            }
        }
    if borders:
        result['borders'] = borders
    if cell.number_format != 'General':
        if cell.number_format not in ('yyyy-mm-dd hh:mm:ss',
                                      'mmmm" "d", "yyyy', '#,##0',
                                      ACCOUNTING_FORMAT, COMPACT_FORMAT):
            raise ValueError('Unsupported source number format: ' +
                             cell.number_format)
        result['numberFormat'] = {
            'type':
            'DATE_TIME' if cell.number_format == 'yyyy-mm-dd hh:mm:ss' else
            'DATE' if cell.number_format == 'mmmm" "d", "yyyy' else 'NUMBER',
            'pattern':
            cell.number_format
        }
    return result


def remove_source(task):
    """Take upstream's template out of the agent's Drive once it is copied.

    Upstream's preprocess copies a Sheet another account shares by link,
    which the agent's own Drive listing never shows. `seed` creates it in
    the agent's Drive, under its upstream ID, only so that copy can find
    it, and the Drive fake lists every file it holds, so the model found
    two Sheets of the same name where upstream's finds one. Nothing reads
    it after the copy has been checked: no grader, task text or doc names
    its ID.

    Args:
        task (str): the task name.
    """
    source = json.loads((source_folder(task) / 'source.json').read_text())
    gws.call('DELETE', f"/drive/v3/files/{source['sheet_id']}")


def validate_seed(box, task):
    """Fail when an upstream seeder swallowed a Sheets setup error."""
    if task == 'woocommerce-stock-alert':
        sheet_id = box.files.read('/opt/toolathlon/tasks/finalpool/' + task +
                                  '/files/sheet_id.txt').strip()
    else:
        config = json.loads(box.files.read('/work/config.json'))
        sheet_id = config['spreadsheet_id']
    source = json.loads((source_folder(task) / 'source.json').read_text())
    if not sheet_id or sheet_id == source['sheet_id']:
        raise ValueError('Original task did not copy its source Sheet')
    original = gws.grid(source['sheet_id'])
    copied = gws.grid(sheet_id)
    if copied['sheets'] != original['sheets']:
        raise ValueError(
            'Copied Sheet values or formatting differ from source')


# Mirage's own source under `integ/` is recorded in the receipt but not
# compared. It is this project's code rather than a foreign pin, and one
# commit in five touches the fake it holds, so requiring byte-identity
# blocked a restore every time the fake was edited -- a re-capture demanded
# by a refactor, not by the fixture having moved. The hashes stay in the
# receipt because `mirage_source` reads the folder: they say which bytes
# actually served the capture. What the restore is checked against instead
# is everything this repo controls and everything the capture itself is of.
MIRAGE_SOURCE = ('submodule/mirage/integ/server/',
                 'submodule/mirage/integ/prisma/',
                 'submodule/mirage/integ/fixtures/')


def same_sources(captured, current):
    # Fixture adapters can change while a restored API state stays identical,
    # so the adapters' own hashes stay in the comparison: the original
    # workbook, the npm dependency closure the fake runs on (recorded under
    # the submodule's name, read from `docker/`), this repo's fixture
    # adapters, and the redirect.
    if not captured:
        return False

    def checked(record):
        return {
            name: sha
            for name, sha in record['inputs_sha256'].items()
            if not name.startswith(MIRAGE_SOURCE)
        }

    return (captured['source'] == current['source']
            and checked(captured) == checked(current)
            and captured['redirect_sha256'] == current['redirect_sha256'])
